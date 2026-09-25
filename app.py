"""
AI-Powered Resume Screener & Talent Matcher
=============================================
A hybrid-AI Streamlit app:
"""

import io
import json
import logging
import os
import re

import numpy as np
import pandas as pd
import streamlit as st

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("resume_screener")

# ===========================================================================
# CONFIG
# ===========================================================================
APP_NAME = "Super TalentMatch AI"
APP_TAGLINE = "Unified Resume Extraction & Deep LLM Screening"
EMBEDDING_MODEL_NAME = "all-MiniLM-L6-v2"
GROQ_MODEL = "openai/gpt-oss-120b"  # Updated to the requested model
ACCEPTED_TYPES = ["pdf", "docx"]
DEFAULT_TOP_N = 3
MAX_TOP_N = 10

# ===========================================================================
# PAGE CONFIG + CUSTOM CSS
# ===========================================================================
st.set_page_config(
    page_title=f"{APP_NAME} | HR Dashboard",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded",
)

CUSTOM_CSS = """
<style>
html, body, [class*="css"] {
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto,
                 "Helvetica Neue", Arial, sans-serif;
    color: #334155; 
}
.stApp { background: #F8FAFC; }

/* ---- Header ---- */
.tm-header {
    padding: 1.6rem 2rem;
    border-radius: 16px;
    margin-bottom: 1.4rem;
    background: linear-gradient(135deg, #0F172A 0%, #1E293B 100%);
}
.tm-header h1 {
    color: #FFFFFF;
    font-size: 2.1rem;
    font-weight: 800;
    margin: 0;
    letter-spacing: -0.01em;
}
.tm-header p {
    color: #94A3B8;
    margin-top: 0.3rem;
    margin-bottom: 0;
    font-size: 1rem;
}

/* ---- Cards ---- */
.tm-card {
    background: #FFFFFF;
    border: 1px solid #E2E8F0;
    border-radius: 14px;
    padding: 1.3rem 1.5rem;
    margin-bottom: 1.1rem;
    box-shadow: 0 1px 4px rgba(15, 23, 42, 0.04);
}
.tm-card h4 { margin-top: 0; color: #1E293B; }

/* ---- Metric-style score chip ---- */
.tm-metric-box {
    background: #F1F5F9;
    border-radius: 10px;
    padding: 0.7rem 0.9rem;
    text-align: center;
}
.tm-metric-box .tm-value { font-size: 1.6rem; font-weight: 800; }
.tm-metric-box .tm-label { font-size: 0.78rem; color: #64748B; text-transform: uppercase; letter-spacing: 0.04em; }
.tm-score-high { color: #15803D; }
.tm-score-mid { color: #B45309; }
.tm-score-low { color: #B91C1C; }

.tm-tier-badge {
    display: inline-block;
    font-size: 0.72rem;
    font-weight: 700;
    padding: 0.15rem 0.55rem;
    border-radius: 999px;
    background: #EEF2FF;
    color: #4338CA;
    margin-left: 0.5rem;
    vertical-align: middle;
}

/* ---- Sidebar ---- */
section[data-testid="stSidebar"] {
    background: #0F172A;
}
section[data-testid="stSidebar"] * { color: #E2E8F0 !important; }
section[data-testid="stSidebar"] input { color: #1E293B !important; }

/* ---- Buttons ---- */
.stButton>button[kind="primary"] {
    background: #4F46E5;
    border: none;
    color: #fff;
    font-weight: 700;
    border-radius: 10px;
}
.stButton>button[kind="primary"]:hover { background: #4338CA; }
.stDownloadButton>button {
    border-radius: 10px;
    border: 1.5px solid #4F46E5;
    color: #4F46E5;
    font-weight: 700;
}
.stDownloadButton>button:hover { background: #4F46E5; color: #fff; }

.block-container { padding-top: 1.8rem; }
</style>
"""
st.markdown(CUSTOM_CSS, unsafe_allow_html=True)

def render_header():
    st.markdown(
        f"""
        <div class="tm-header">
            <h1>🎯 {APP_NAME}</h1>
            <p>{APP_TAGLINE}</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

# ===========================================================================
# FILE TEXT EXTRACTION
# ===========================================================================
def extract_text_from_pdf(file_bytes: bytes) -> str:
    import pdfplumber
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                text_parts.append(page_text)
    return "\n".join(text_parts)

def extract_text_from_docx(file_bytes: bytes) -> str:
    import docx
    document = docx.Document(io.BytesIO(file_bytes))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)

def extract_resume_text(uploaded_file):
    name = uploaded_file.name.lower()
    ext = name.rsplit(".", 1)[-1] if "." in name else ""
    try:
        file_bytes = uploaded_file.read()
        if ext == "pdf":
            text = extract_text_from_pdf(file_bytes)
        elif ext == "docx":
            text = extract_text_from_docx(file_bytes)
        else:
            st.error(f"⚠️ Unsupported file type for **{uploaded_file.name}** (.{ext}) — skipped.")
            return None

        if not text.strip():
            st.error(f"⚠️ No readable text found in **{uploaded_file.name}** — skipped.")
            return None
        return text
    except Exception as exc:
        st.error(f"⚠️ Could not read **{uploaded_file.name}**: {exc} — skipped.")
        return None

def guess_candidate_name(resume_text: str, fallback_filename: str) -> str:
    for line in resume_text.splitlines()[:5]:
        candidate = line.strip()
        if not candidate:
            continue
        words = candidate.split()
        alpha_ratio = sum(c.isalpha() or c.isspace() for c in candidate) / max(len(candidate), 1)
        if 1 <= len(words) <= 4 and alpha_ratio > 0.85 and "@" not in candidate:
            return candidate.title()
    return re.sub(r"\.(pdf|docx)$", "", fallback_filename, flags=re.IGNORECASE)

# ===========================================================================
# PHASE 1 — LOCAL EMBEDDING-BASED SCREENING 
# ===========================================================================
@st.cache_resource(show_spinner="Loading local embedding model...")
def load_embedding_model():
    from sentence_transformers import SentenceTransformer
    return SentenceTransformer(EMBEDDING_MODEL_NAME)

def compute_match_scores(jd_text: str, resume_texts: list) -> list:
    from sklearn.metrics.pairwise import cosine_similarity
    model = load_embedding_model()
    all_texts = [jd_text] + resume_texts
    embeddings = model.encode(all_texts, show_progress_bar=False)

    jd_embedding = embeddings[0].reshape(1, -1)
    resume_embeddings = embeddings[1:]

    similarities = cosine_similarity(jd_embedding, resume_embeddings)[0]
    scores = np.clip(similarities, 0, 1) * 100
    return [round(float(s), 1) for s in scores]

# ===========================================================================
# PHASE 2 — DEEP AI ANALYSIS VIA GROQ 
# ===========================================================================
def get_groq_api_key():
    try:
        if "GROQ_API_KEY" in st.secrets:
            return st.secrets["GROQ_API_KEY"]
    except Exception:
        pass
    return os.environ.get("GROQ_API_KEY")

@st.cache_resource(show_spinner=False)
def get_groq_client(api_key: str):
    from groq import Groq
    return Groq(api_key=api_key)

def build_deep_analysis_prompt(resume_text: str, jd_text: str) -> str:
    trimmed_resume = resume_text[:12000]
    return f"""You are an expert technical recruiter and ATS parser. Compare this CANDIDATE RESUME against the JOB DESCRIPTION and extract the candidate's details.

CANDIDATE RESUME:
{trimmed_resume}

JOB DESCRIPTION:
{jd_text}

Return ONLY a JSON object with exactly these keys:
- "email": The candidate's email address (if found, else "N/A").
- "education": The highest degree or qualification (e.g., "Bachelors in HR" or "N/A").
- "university_name": The name of the university or institution (if found, else "N/A").
- "experience_years": Total years of experience (e.g., "3 Years", "5+ Years", or "N/A").
- "latest_experience": Their most recent job title and company (if found, else "N/A").
- "custom_questions": a list of EXACTLY 3 highly specific interview questions, each targeting a genuine gap or weakness you identified between this candidate's resume and the job description.

Do not wrap the JSON in markdown code fences. Do not add any commentary before or after the JSON.
"""

def analyze_with_groq(client, resume_text: str, jd_text: str, file_name: str):
    try:
        prompt = build_deep_analysis_prompt(resume_text, jd_text)
        response = client.chat.completions.create(
            model=GROQ_MODEL,
            messages=[{"role": "user", "content": prompt}],
            response_format={"type": "json_object"},
            temperature=0.3,
        )
        result = json.loads(response.choices[0].message.content)
        
        # Ensure keys exist
        result.setdefault("email", "N/A")
        result.setdefault("education", "N/A")
        result.setdefault("university_name", "N/A")
        result.setdefault("experience_years", "N/A")
        result.setdefault("latest_experience", "N/A")
        result.setdefault("custom_questions", [])
        return result
    except Exception as exc:
        st.error(f"⚠️ Groq deep analysis failed for **{file_name}**: {exc}")
        return {
            "email": "N/A", "education": "N/A", "university_name": "N/A",
            "experience_years": "N/A", "latest_experience": "N/A",
            "custom_questions": []
        }

# ===========================================================================
# EXCEL EXPORT
# ===========================================================================
def build_results_dataframe(results: list) -> pd.DataFrame:
    rows = []
    for r in results:
        rows.append({
            "Name": r["name"],
            "Match Score (%)": r["match_score"],
            "Email": r.get("email", "N/A"),
            "Education": r.get("education", "N/A"),
            "University Name": r.get("university_name", "N/A"),
            "Experience (Years)": r.get("experience_years", "N/A"),
            "Latest Experience": r.get("latest_experience", "N/A"),
            "Custom Interview Questions": " | ".join(r.get("custom_questions", [])) or "—",
        })
    df = pd.DataFrame(rows).sort_values("Match Score (%)", ascending=False).reset_index(drop=True)
    return df

def dataframe_to_formatted_excel_bytes(df: pd.DataFrame) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="xlsxwriter") as writer:
        df.to_excel(writer, index=False, sheet_name="Candidates")
        workbook = writer.book
        worksheet = writer.sheets["Candidates"]

        header_format = workbook.add_format({
            "bold": True, "bg_color": "#1E293B", "font_color": "#FFFFFF",
            "border": 1, "align": "center", "valign": "vcenter",
        })
        wrap_format = workbook.add_format({"text_wrap": True, "valign": "top"})

        for col_idx, col_name in enumerate(df.columns):
            worksheet.write(0, col_idx, col_name, header_format)
            longest_value = df[col_name].astype(str).map(len).max() if len(df) else 0
            width = max(longest_value, len(col_name)) + 3
            width = min(width, 55) 
            worksheet.set_column(col_idx, col_idx, width, wrap_format if col_name in
                                  ("Custom Interview Questions", "Latest Experience") else None)
        worksheet.freeze_panes(1, 0)

    buffer.seek(0)
    return buffer.getvalue()

# ===========================================================================
# SCORE DISPLAY HELPERS
# ===========================================================================
def score_class(score: float) -> str:
    if score >= 75: return "tm-score-high"
    if score >= 50: return "tm-score-mid"
    return "tm-score-low"

def render_score_box(score: float):
    st.markdown(
        f'<div class="tm-metric-box">'
        f'<div class="tm-value {score_class(score)}">{score}%</div>'
        f'<div class="tm-label">Match Score</div>'
        f'</div>',
        unsafe_allow_html=True,
    )

# ===========================================================================
# SESSION STATE
# ===========================================================================
if "results" not in st.session_state:
    st.session_state.results = []

# ===========================================================================
# SIDEBAR
# ===========================================================================
with st.sidebar:
    st.markdown(f"## 🎯 {APP_NAME}")
    st.caption(APP_TAGLINE)
    st.markdown("---")

    st.markdown("**Groq API Key**")
    prefilled_key = get_groq_api_key() or ""
    groq_api_key = st.text_input(
        "Groq API Key", value=prefilled_key, type="password",
        label_visibility="collapsed", placeholder="gsk_...",
        help="Get a free key at console.groq.com/keys",
    )
    if groq_api_key:
        st.success("Groq API key loaded.")
    else:
        st.warning("Add a Groq API key to enable data extraction and deep AI analysis.")

    st.markdown("---")
    top_n = st.slider(
        "Top candidates for deep AI analysis", min_value=1, max_value=MAX_TOP_N,
        value=DEFAULT_TOP_N,
    )

    st.markdown("---")
    if st.button("🗑️ Clear Results", use_container_width=True):
        st.session_state.results = []
        st.rerun()

# ===========================================================================
# MAIN PAGE
# ===========================================================================
render_header()

col_jd, col_upload = st.columns(2, gap="large")

with col_jd:
    st.markdown('<div class="tm-card">', unsafe_allow_html=True)
    st.markdown("#### 📋 Job Description")
    jd_text = st.text_area(
        "Paste the job description", height=260, label_visibility="collapsed",
        placeholder="Paste the full job description here...",
    )
    st.markdown("</div>", unsafe_allow_html=True)

with col_upload:
    st.markdown('<div class="tm-card">', unsafe_allow_html=True)
    st.markdown("#### 📥 Upload Resumes")
    st.caption("Drag and drop files here. Supported: PDF, DOCX.")
    uploaded_files = st.file_uploader(
        "Upload resumes", type=ACCEPTED_TYPES, accept_multiple_files=True,
        label_visibility="collapsed",
    )
    if uploaded_files:
        st.caption(f"{len(uploaded_files)} file(s) ready.")
    st.markdown("</div>", unsafe_allow_html=True)

screen_clicked = st.button(
    "🚀 Process & Screen Candidates", type="primary", use_container_width=True,
    disabled=not uploaded_files or not jd_text.strip(),
)

# ---------------------------------------------------------------------------
# PIPELINE
# ---------------------------------------------------------------------------
if screen_clicked and uploaded_files and jd_text.strip():
    with st.spinner("Extracting text from resumes..."):
        extracted = []
        for f in uploaded_files:
            text = extract_resume_text(f)
            if text:
                extracted.append((f.name, text))

    if not extracted:
        st.error("No resumes could be read.")
    else:
        with st.spinner("Locally scoring resumes against the JD..."):
            resume_texts = [text for _, text in extracted]
            scores = compute_match_scores(jd_text, resume_texts)

        candidates = [
            {
                "file_name": file_name,
                "name": guess_candidate_name(text, file_name),
                "text": text,
                "match_score": score,
                "analyzed": False,
            }
            for (file_name, text), score in zip(extracted, scores)
        ]
        candidates.sort(key=lambda c: c["match_score"], reverse=True)

        if groq_api_key:
            top_candidates = candidates[:top_n]
            client = get_groq_client(groq_api_key)
            progress = st.progress(0.0, text="Starting deep AI data extraction...")
            for i, cand in enumerate(top_candidates):
                progress.progress(i / len(top_candidates), text=f"Analyzing {cand['name']} ...")
                analysis = analyze_with_groq(client, cand["text"], jd_text, cand["file_name"])
                cand["analyzed"] = True
                cand["email"] = analysis["email"]
                cand["education"] = analysis["education"]
                cand["university_name"] = analysis["university_name"]
                cand["experience_years"] = analysis["experience_years"]
                cand["latest_experience"] = analysis["latest_experience"]
                cand["custom_questions"] = analysis["custom_questions"]
                progress.progress((i + 1) / len(top_candidates))
            progress.empty()
        
        st.session_state.results = candidates
        st.success(f"Successfully processed {len(candidates)} candidates!")

# ---------------------------------------------------------------------------
# RESULTS
# ---------------------------------------------------------------------------
if st.session_state.results:
    results = st.session_state.results

    st.markdown('<div class="tm-card">', unsafe_allow_html=True)
    st.markdown("#### 📊 Screening Overview")
    m1, m2 = st.columns(2)
    m1.metric("Total Candidates Processed", len(results))
    avg_score = round(sum(r["match_score"] for r in results) / len(results), 1)
    m2.metric("Average Match Score", f"{avg_score}%")
    st.markdown("</div>", unsafe_allow_html=True)

    st.markdown('<div class="tm-card">', unsafe_allow_html=True)
    st.markdown("#### 🧾 Candidate Details & Rankings")
    for rank, cand in enumerate(results, start=1):
        header_html = f"#{rank} — {cand['name']} | Score: {cand['match_score']}%"
        with st.expander(header_html, expanded=(rank <= 3)):
            score_col, detail_col = st.columns([1.5, 3.5])
            with score_col:
                render_score_box(cand["match_score"])
                if cand["analyzed"]:
                    st.markdown('<br><span class="tm-tier-badge">Extracted Data</span>', unsafe_allow_html=True)
            with detail_col:
                if cand["analyzed"]:
                    st.markdown(f"**📧 Email:** {cand.get('email', 'N/A')}")
                    st.markdown(f"**🎓 Education:** {cand.get('education', 'N/A')}")
                    st.markdown(f"**🏛️ University:** {cand.get('university_name', 'N/A')}")
                    st.markdown(f"**⏳ Experience (Years):** {cand.get('experience_years', 'N/A')}")
                    st.markdown(f"**💼 Latest Role:** {cand.get('latest_experience', 'N/A')}")
                    st.markdown("---")
                    st.markdown("**🗣️ Suggested Interview Questions:**")
                    for q in cand.get("custom_questions", []):
                        st.markdown(f"- {q}")
                else:
                    st.caption("Not selected for AI data extraction (outside the top-N cutoff).")
    st.markdown("</div>", unsafe_allow_html=True)

    st.markdown('<div class="tm-card">', unsafe_allow_html=True)
    st.markdown("#### ⬇️ Export Master Sheet")
    df = build_results_dataframe(results)
    st.download_button(
        "Download Formatted Excel (.xlsx)",
        data=dataframe_to_formatted_excel_bytes(df),
        file_name="TalentMatch_Candidates.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )
    st.dataframe(df, use_container_width=True)
    st.markdown("</div>", unsafe_allow_html=True)

else:
    st.info("Paste a Job Description, upload resumes, and click **Process & Screen Candidates** to get started.")
